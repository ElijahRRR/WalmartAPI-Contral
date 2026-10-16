"""weight_fix — 按名单把已上架商品的发货重量改成固定值(临时工具,不进调度;危险:缺省即真跑,空跑用 --dry-run)。

用途(所有者 2026-10-16):线上有一批发货重量明显不对的品(Item Report 里 1814 磅、
13200 磅这类),来自 9/7 以前的旧解析、旧系统和跟卖 D 列。所有者整理好名单,
本工具把名单里的品的发货重量**一次性**改成同一个值(缺省 11 磅)。

  python cli.py weight_fix -p file=<名单.xlsx> --dry-run                 # 先看:读到几行、对上几条、跳过几条
  python cli.py weight_fix -p file=<名单.xlsx> -p store=谭总2 -p limit=1  # 先真发 1 条试水
  python cli.py weight_fix -p file=<名单.xlsx>                            # 全量
  -p weight=11       改成多少磅,缺省 11
  -p sheet=<表名>    xlsx 默认第一张工作表

**名单**:xlsx 或 csv,表头里要有「店铺」和「SKU」两列(也认 store / sku),其余列
一概不看。同一 (店铺, SKU) 出现多次只发一次。

**通道:MP_MAINTENANCE 部分更新**,每条载荷只有三样:sku、productIdentifiers、
ShippingWeight。依据是官方规范原件(摘录与出处见
`refdata/specs/MP_MAINTENANCE_5.0.20260703-18_22_27-api.orderable_excerpt.json`):
Orderable 20 个属性里有 `ShippingWeight`(Shipping Weight (lbs),number),必填只有
sku / productIdentifiers,其余不带就不动 —— **价格、库存、标题都不碰**。
productIdentifiers 取库里 `catalog.walmart_items.upc`,与维护链改标题同一个形态
(`maintenance_intents.build_title_item`:UPC)。不用 MP_ITEM_MATCH:那是整条覆盖,
载荷里每个字段都会被重写,为了改一个重量要把现价、库存都原样带回去。

**哪些行不发**(摘要里逐类计数):
  · 目录里查不到这个 (店铺, SKU),或已缺席(missing_since 非空):改了也没意义;
  · 库里这条没有 UPC:载荷必填,拿不到不猜;
  · 店铺不在可调用列表里(未启用 / 没代理 / 没凭证)。

**请求没发出 ≠ 被拒**(2026-10-16 生产实见:A085朱丽霖 代理 TLS 握手 EOF):取 token /
代理阶段失败时请求根本没到沃尔玛,台账落 failed、可原样重占。这种店按店维标准
**串行补试一遍**,仍失败就在摘要首行点名「N 条请求没发出」,修好代理直接重跑即可。
沃尔玛明确拒收(4xx)不补试,摘要写「沃尔玛拒收」。

⚠ **回执 SUCCESS 不等于改成了**。MP_MAINTENANCE 有过"回执成功、内容静默丢弃"的
生产实证(2026-09-04 带 SkuUpdate 的改码 feed,那是规范外字段)。ShippingWeight 在
规范内,但本仓从没用这个通道改过重量,所以**第一次真跑先 `-p limit=1`**,等回执
成功后到后台或 Item Report 看那一条的 Shipping Weight 变了,再跑全量。

⚠ 只改线上值,**不改上架规则**:上架链与改码链仍按 `mp_mapper.shipping_weight_ex`
写重量(> 11 磅或读不出写 1 磅)。名单里的品以后若被改码,重量会按那条规则重写。

回执由 feed_poll 统一轮询落 `ops.feed_items`(workflow=weight_fix);不进维护记录表
(`maint_sheet` 只读 workflow=maintenance),也不进产品病历(`product_events.receipt_in_ledger`
只收反补来源的 MP_MAINTENANCE)。同一批载荷还在处理中时重跑,feed 台账防重拦下
(`dedup`),不会重复提交。「实际结果」复核里这些明细没有处置建议行,记为「无目标值」;
它把同一 (店, SKU) 后发的 MP_MAINTENANCE 当作覆盖了前一个,所以同 SKU 此前的改标题
明细不再被判(标题维护 2026-09-05 起整路停着,实际影响为零)。
"""

import csv
import logging

from api import feeds
from registry import db
from services import store_retry, stores as stores_svc

DANGEROUS = True
SUPPORTS_STORE = True

logger = logging.getLogger("workflows.weight_fix")

_FEED_TYPE = "MP_MAINTENANCE"
DEFAULT_WEIGHT_LBS = 11.0       # 所有者 2026-10-16:名单里的品统一改成 11 磅

_STORE_HEADERS = ("店铺", "store")
_SKU_HEADERS = ("sku",)         # 表头比较一律小写,「SKU」也命中

# 名单里 (店铺, SKU) 在目录里的现状。只认还在目录里的行(missing_since IS NULL)。
# ⚠ 参数带显式类型(本仓 SQL 纪律:PG 推不出类型在生产上连炸过三次)。
_SQL_ITEMS = """
SELECT w.sku, w.upc
FROM catalog.walmart_items w
WHERE w.store = %(store)s::text
  AND w.sku = ANY(%(skus)s::text[])
  AND w.missing_since IS NULL
"""


def _cell(v) -> str:
    """输入:表格单元格值 → 输出:去空白的字符串(整数形态的浮点去掉 .0)。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _header_index(header: list, names: tuple) -> int | None:
    """输入:表头行 + 认可的列名 → 输出:列下标或 None。"""
    norm = [_cell(h).lower() for h in header]
    for n in names:
        if n.lower() in norm:
            return norm.index(n.lower())
    return None


def read_list(path: str, sheet: str = "") -> tuple[dict[str, list[str]], dict]:
    """输入:名单文件路径(xlsx/csv)+ 可选工作表名 → 输出:({店铺: [sku…]}, 计数)。

    只取「店铺」「SKU」两列;空行跳过,同一 (店铺, SKU) 只留一次,保持文件里的先后。
    表头缺列直接抛 —— 读错列比报错更坏(会把别的列当 SKU 发出去)。
    """
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        header, body = (rows[0] if rows else []), rows[1:]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet and sheet not in wb.sheetnames:
            raise LookupError(f"工作表「{sheet}」不存在,本文件有:{wb.sheetnames}")
        ws = wb[sheet] if sheet else wb.worksheets[0]
        # 有些导出文件把 dimension 声明成单格,read_only 只读出 1 行(kpi 2026-08-06 实证)
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        it = ws.iter_rows(values_only=True)
        header = list(next(it, []) or [])
        body = it
    i_store = _header_index(header, _STORE_HEADERS)
    i_sku = _header_index(header, _SKU_HEADERS)
    if i_store is None or i_sku is None:
        raise ValueError(f"名单表头里要有「店铺」和「SKU」两列,实际表头:"
                         f"{[_cell(h) for h in header]}")
    by_store: dict[str, list[str]] = {}
    seen: set = set()
    stat = {"rows": 0, "blank": 0, "dup": 0}
    for row in body:
        row = list(row or [])
        store = _cell(row[i_store]) if i_store < len(row) else ""
        sku = _cell(row[i_sku]) if i_sku < len(row) else ""
        if not store and not sku:
            continue                       # 整行空(表尾常见),不算数
        stat["rows"] += 1
        if not store or not sku:
            stat["blank"] += 1
            continue
        if (store, sku) in seen:
            stat["dup"] += 1
            continue
        seen.add((store, sku))
        by_store.setdefault(store, []).append(sku)
    return by_store, stat


class _NotSent(Exception):
    """本店有切片请求没发出(取 token / 代理阶段失败)→ 交 fan_out 串行补试一遍。"""


def build_item(sku: str, upc: str, weight: float) -> dict:
    """输入:SKU + UPC + 磅数 → 输出:MP_MAINTENANCE 只改发货重量的 MPItem。

    键集合由 tests/test_weight_fix.py 对照官方规范原件摘录校验。
    """
    return {"Orderable": {"sku": str(sku),
                          "productIdentifiers": {"productId": str(upc),
                                                 "productIdType": "UPC"},
                          "ShippingWeight": float(weight)}}


def _fix_store(store: dict, skus: list[str], weight: float, limit: int,
               preview: bool, not_sent: dict | None = None) -> dict:
    """输入:店铺 + 名单里该店的 SKU + 磅数 + 单店上限 + 是否空跑 → 输出:{name, lines, 各计数}。

    读库失败或提交时凭证死直接抛,交 fan_out 按店维标准处理(失败店串行补试一遍;
    补试时同一批载荷若已提交成功,feed 台账防重会拦成 dedup,不会重复提交)。
    请求**没发出**的切片(取 token / 代理阶段失败,submit_feed 标 retryable)也抛
    `_NotSent` 交补试:那一片在台账里是 failed、可原样重占,补交的是同一载荷。
    `not_sent` 记下哪家店是这种失败,补试仍失败时摘要照实写「请求没发出」。
    """
    name = store["name"]
    with db.pg_conn() as conn, conn.cursor() as cur:
        cur.execute(_SQL_ITEMS, {"store": name, "skus": skus})
        found = {str(s): (str(u).strip() if u else "") for s, u in cur.fetchall()}
    not_found = [s for s in skus if s not in found]
    no_upc = [s for s in skus if s in found and not found[s]]
    ready = [s for s in skus if found.get(s)]
    held = 0
    if limit and len(ready) > limit:
        held = len(ready) - limit
        ready = ready[:limit]
    res = {"name": name, "lines": [], "listed": len(skus), "planned": len(ready),
           "not_found": len(not_found), "no_upc": len(no_upc), "held": held,
           "submitted": 0, "dedup": 0, "failed": 0, "unknown": 0, "not_sent": 0,
           "feeds": []}
    line = f"  {name}:名单 {len(skus)} 条,要改 {len(ready)} 条"
    if held:
        line += f"(limit 截下 {held} 条,下次再发)"
    if not_found:
        line += (f";目录里没有或已缺席 {len(not_found)} 条(如 "
                 f"{','.join(not_found[:3])})")
    if no_upc:
        line += f";库里没 UPC {len(no_upc)} 条(如 {','.join(no_upc[:3])})"
    res["lines"].append(line)
    if not ready or preview:
        return res
    entries = [build_item(s, found[s], weight) for s in ready]
    for r in feeds.submit_feed(store, _FEED_TYPE, entries, workflow="weight_fix"):
        # submit_feed 的结局:submitted / dedup / failed / unknown
        # (不传 defer_settle 就不会有 deferred;真出现了也按不确定报)
        if r["outcome"] == "failed" and r.get("retryable"):
            outcome = "not_sent"        # 取 token / 代理阶段失败:请求没发出,不是被拒
        elif r["outcome"] in ("submitted", "dedup", "failed"):
            outcome = r["outcome"]
        else:
            outcome = "unknown"
        res[outcome] += r["count"]
        if r.get("feed_id"):
            res["feeds"].append(r["feed_id"])
    if not_sent is not None:
        if res["not_sent"]:
            not_sent[name] = res["not_sent"]
        else:
            not_sent.pop(name, None)    # 补试成功:撤掉第一轮的记号
    if res["not_sent"]:
        raise _NotSent(f"{name}:{res['not_sent']} 条请求没发出(取 token / 代理阶段失败)")
    parts = [f"已提交 {res['submitted']} 条"]
    if res["dedup"]:
        parts.append(f"同一批还在处理中、防重拦下 {res['dedup']} 条")
    if res["failed"]:
        parts.append(f"⚠ 沃尔玛拒收 {res['failed']} 条(看日志里的 HTTP 响应)")
    if res["unknown"]:
        parts.append(f"⚠ 结局不确定 {res['unknown']} 条"
                     "(feed_poll 的 pending 对账会接手,不要手工补发)")
    res["lines"].append("    " + ",".join(parts)
                        + (f";feed:{','.join(res['feeds'])}" if res["feeds"] else ""))
    return res


def _weight(v) -> float:
    """输入:-p weight= 原值 → 输出:磅数;非正数/读不出直接抛(宁炸不发错值)。"""
    if v in (None, ""):
        return DEFAULT_WEIGHT_LBS
    w = float(v)
    if not w > 0:
        raise ValueError(f"weight 必须是正数磅数,收到 {v!r}")
    return round(w, 2)


def run(params: dict) -> str:
    """输入:params(file 必填;store / limit / weight / sheet 可选)→ 输出:改重量摘要。"""
    path = str(params.get("file") or "").strip()
    if not path:
        return "⚠ 缺名单:-p file=<名单.xlsx 或 .csv>(表头要有「店铺」「SKU」)"
    weight = _weight(params.get("weight"))
    limit = int(params.get("limit") or 0)
    only = str(params.get("store") or "").strip()
    preview = bool(params.get("dry_run"))
    mode = "[DRY-RUN] " if preview else ""

    by_store, stat = read_list(path, str(params.get("sheet") or "").strip())
    if only:
        by_store = {k: v for k, v in by_store.items() if k == only}
        if not by_store:
            return f"⚠ 名单里没有店铺 {only} 的行"
    names = sorted(by_store, key=stores_svc.sort_key)
    store_list = stores_svc.load_stores(filter_names=names) if names else []
    callable_names = {s["name"] for s in store_list}
    uncallable = [n for n in names if n not in callable_names]
    head_skip = ""
    if uncallable:
        n_rows = sum(len(by_store[n]) for n in uncallable)
        head_skip = (f";⚠ 不在可调用店铺里(未启用/没代理/没凭证)整店跳过 "
                     f"{len(uncallable)} 店 {n_rows} 条:{','.join(uncallable)}")
    read_note = (f"名单读到 {stat['rows']} 行,{len(by_store)} 店"
                 + (f",缺店铺或 SKU {stat['blank']} 行" if stat["blank"] else "")
                 + (f",重复 {stat['dup']} 行" if stat["dup"] else ""))
    if not store_list:
        return (f"{mode}改发货重量 → {weight:g} 磅:没有可调用的店,无事可做;"
                + read_note + head_skip)

    not_sent: dict[str, int] = {}

    def _one(store: dict) -> dict:
        return _fix_store(store, by_store[store["name"]], weight, limit, preview,
                          not_sent)

    results, dead, absent, gate_note = store_retry.fan_out(
        store_list, _one, stores_svc.STORE_WORKERS, log_label="改发货重量")
    results.sort(key=lambda r: stores_svc.sort_key(r["name"]))
    sums = {k: sum(r[k] for r in results)
            for k in ("planned", "not_found", "no_upc", "held",
                      "submitted", "dedup", "failed", "unknown")}
    missed = [f"{s}(凭证失效)" for s in sorted(dead)] + \
        [f"{s}({w}" + (f",{not_sent[s]} 条请求没发出:取 token / 代理阶段失败"
                       if s in not_sent else "") + ")"
         for s, w in sorted(absent)]
    if preview:
        tail = "(dry-run:一条都没发)"
    else:
        tail = (f",已提交 {sums['submitted']} 条,结果由 feed_poll 回写"
                + (f",防重拦下 {sums['dedup']}" if sums["dedup"] else "")
                + (f",⚠ 沃尔玛拒收 {sums['failed']}" if sums["failed"] else "")
                + (f",⚠ 结局不确定 {sums['unknown']}" if sums["unknown"] else ""))
    skipped = []
    if sums["not_found"]:
        skipped.append(f"目录里没有或已缺席 {sums['not_found']}")
    if sums["no_upc"]:
        skipped.append(f"没 UPC {sums['no_upc']}")
    if sums["held"]:
        skipped.append(f"limit 截下 {sums['held']}")
    head = (f"{mode}改发货重量 → {weight:g} 磅:{len(results)} 店,要改 "
            f"{sums['planned']} 条{tail}"
            + (f";不发:{','.join(skipped)}" if skipped else "")
            + (f";⚠ 失败 {len(missed)} 店(已串行补试一遍):{','.join(missed)}"
               "(修好后直接重跑即可,没发出的会重发,已提交的被防重挡住、不会重复提交)"
               if missed else "")
            + head_skip)
    lines = [head, f"  {read_note}"]
    if gate_note:
        lines.append(f"  {gate_note}")
    for r in results:
        lines.extend(r["lines"])
    return "\n".join(lines)
